import io

import pandas as pd
from openpyxl import load_workbook

from database import delete_all_data, summary_stats, upsert_dataframe
from excel_export import build_workbook


def _sample_dataframe():
    return pd.DataFrame(
        [
            {
                "FIR NUMBER": "FIR-001",
                "REGION": "NORTH",
                "STATE": "DELHI",
                "TERRITORY": "T1",
                "FINANCIAL_YEAR": "FY 2024-25",
                "MONTH_NAME": "April",
                "WEEKDAY_NAME": "Monday",
                "DAY_OF_MONTH": 1,
                "TIME_BUCKET": "05:01 to 10:00",
                "IS_INCIDENT": 1,
                "FATALITIES_TOTAL": 1,
                "INJURIES_TOTAL": 2,
                "FIR_TURNAROUND_DAYS": 3,
                "DIR STATUS": "CLOSED",
                "ROOT CAUSE-1": "Equipment Failure",
                "ROOT-CAUSE-2": "Human Error",
                "ROOT-CAUSE-3": None,
                "CONSUMER TYPE": "DOMESTIC",
                "ZONE": "ZONE A",
            }
        ]
    )


def test_delete_all_data_removes_saved_records(tmp_path):
    db_path = tmp_path / "incident_data.db"
    df = _sample_dataframe()
    upsert_dataframe(df, db_path)
    assert summary_stats(db_path)["total_incidents"] == 1

    delete_all_data(db_path)

    assert summary_stats(db_path)["total_incidents"] == 0


def test_excel_report_has_summary_and_chart_tables(tmp_path):
    df = _sample_dataframe()
    workbook_bytes = build_workbook(df)
    workbook = load_workbook(filename=io.BytesIO(workbook_bytes))

    assert "Summary" in workbook.sheetnames
    assert "Historic Trend" in workbook.sheetnames
    assert "Region Wise" in workbook.sheetnames
    assert "FIR, DIR Status" in workbook.sheetnames

    summary = workbook["Summary"]
    assert summary["A1"].value == "Dashboard Summary"
